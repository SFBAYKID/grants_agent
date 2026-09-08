"""Incident regressions for recent-award filtering, continuation, and exports."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from grant_watch import db, google_sheets
from grant_watch.models import LeadGrade
from grant_watch.slack import conversation, search_recency, tools
from grant_watch.slack.search import search_leads
from grant_watch.slack.search_recency import recent_award_scope
from grant_watch.spreadsheets import search_export_filename
from tests.test_search import _db, _insert

TODAY = date(2026, 9, 8)
BRETT_ASK = (
    "can you put together a spreadsheet of all Pennsylvania and Maine schools "
    "that were awarded a grant recently"
)


@pytest.mark.parametrize(
    "user_text",
    [
        BRETT_ASK,
        "make me a list of New Hampshire schools who were awarded a grant recently",
        "make me a list of Conneticut schools who were awarded a grant recently",
        "make me a list of New York schools who were awarded a grant recently",
        "May I have all recent school grants?",
        "schools that just got funding",
    ],
)
def test_recent_awards_have_inclusive_calendar_bounds(user_text: str) -> None:
    """All includes all matches inside the window, not all historical years."""
    scope = recent_award_scope(user_text, None, today=TODAY)
    assert scope is not None
    assert (scope.start, scope.end) == (date(2026, 3, 8), TODAY)
    assert scope.constrain({"state": "PA", "date_from": "2010-01-01"}) == {
        "state": "PA",
        "record_kind": "award",
        "date_field": "award_received",
        "date_from": "2026-03-08",
        "date_to": "2026-09-08",
    }


@pytest.mark.parametrize(
    ("today", "start"),
    [(date(2026, 8, 31), date(2026, 2, 28)), (date(2024, 8, 31), date(2024, 2, 29))],
)
def test_calendar_window_clamps_short_months(today: date, start: date) -> None:
    """Month ends and leap years must not become fixed 180-day intervals."""
    scope = recent_award_scope(BRETT_ASK, None, today=today)
    assert scope is not None and scope.start == start


@pytest.mark.parametrize(
    "user_text",
    [
        "what about New Hampshire?",
        "Conneticut",
        "NY please",
        "can you put these new york schools in a spreadsheet",
        "all 86 in Excel please",
        "find contacts for the top 5",
    ],
)
def test_state_and_export_followups_inherit_human_scope(user_text: str) -> None:
    """Intervening bot descriptions cannot relax the rep's recent-award window."""
    context = [
        "rep: " + BRETT_ASK,
        "Grant: Here are all years; the most recent is 2025.",
        "rep: what about New Hampshire?",
        "Grant: Try older awards.",
        "rep: CT too",
        "rep: " + user_text,
    ]
    scope = recent_award_scope(user_text, context, today=TODAY)
    assert scope is not None and scope.start == date(2026, 3, 8)


@pytest.mark.parametrize(
    "user_text",
    [
        "all years of school awards",
        "historical awards instead",
        "schools awarded grants in 2019",
        "recent awards in July 2026",
        "recent awards in May",
        "recent awards from the last 30 days",
        "recent awards since 2024-01-01",
        "schools awarded a grant, not recently",
        "schools that have not been awarded a grant recently",
        "don't limit this to recent school awards",
        "recent school awards over the past nine months",
        "recent school awards over the past twenty four months",
        "recent RFPs in Pennsylvania",
        "recently discovered grants",
        "recent spending windows",
        "find California NSGP awards",
        "newest verified award announcements",
    ],
)
def test_explicit_scope_or_new_query_does_not_inherit_default(user_text: str) -> None:
    """Historical, explicit temporal, and different event meanings remain available."""
    assert recent_award_scope(user_text, ["rep: " + BRETT_ASK], today=TODAY) is None


def test_later_human_override_stops_inheritance_but_bot_words_cannot_start_it() -> None:
    """A later human date change supersedes old scope; Grant cannot invent consent."""
    context = ["rep: " + BRETT_ASK, "rep: show historical awards instead"]
    assert recent_award_scope("NY too", context, today=TODAY) is None
    assert (
        recent_award_scope("NY too", ["Grant: Recent awards only"], today=TODAY) is None
    )


def _scripted_client(turns: list[list[dict[str, object]]]) -> type:
    """Create an offline model that proposes omitted or deliberately widened dates."""

    class Messages:
        """Supply each scripted tool turn and a deliberately scope-free final reply."""

        def __init__(self) -> None:
            """Keep execution position private to this fake model session."""
            self.index = 0

        def create(self, **_kwargs: object) -> object:
            """Return model-shaped blocks without making any network request."""
            if self.index < len(turns):
                proposals = turns[self.index]
                self.index += 1
                return SimpleNamespace(
                    stop_reason="tool_use",
                    content=[
                        SimpleNamespace(
                            type="tool_use",
                            name="search_leads",
                            input=args,
                            id=f"{self.index}-{index}",
                        )
                        for index, args in enumerate(proposals)
                    ],
                )
            return SimpleNamespace(
                stop_reason="end_turn",
                content=[
                    SimpleNamespace(
                        type="text",
                        text=json.dumps(
                            {
                                "intent": "question",
                                "reply": "No matching indexed awards.",
                            }
                        ),
                    )
                ],
            )

    class Client:
        """Expose only the model interface used by the conversation."""

        def __init__(self, **_kwargs: object) -> None:
            """Start the fake message stream."""
            self.messages = Messages()

    return Client


def test_two_states_and_zero_result_retry_cannot_escape_scope(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The execution boundary, not model wording, must prevent all-years retries."""
    proposals = [
        [
            {
                "state": state,
                "org_type": "school",
                "export": "excel",
                "result_scope": "all",
            }
            for state in ("PA", "ME")
        ],
        [
            {
                "state": "PA",
                "org_type": "school",
                "date_from": "2010-01-01",
                "date_to": "2030-01-01",
                "limit": 100,
            }
        ],
    ]
    executions: list[dict[str, object]] = []

    def capture(
        _name: str, args: dict[str, object], *_args: object, **_kwargs: object
    ) -> tuple[str, None]:
        """Record actual execution inputs and tempt a wider search with a zero hint."""
        executions.append(args)
        return "No grants matched. Without the date window: 35 matches.", None

    monkeypatch.setattr(search_recency, "utc_today", lambda: TODAY)
    monkeypatch.setattr(conversation, "Anthropic", _scripted_client(proposals))
    monkeypatch.setattr(tools, "run_tool", capture)
    out = conversation.respond(BRETT_ASK, None)
    assert len(executions) == 3
    for args in executions:
        assert args["date_from"] == "2026-03-08"
        assert args["date_to"] == "2026-09-08"
        assert args["date_field"] == "award_received"
        assert args["record_kind"] == "award"
    assert "2026-03-08 through 2026-09-08" in out["reply"]
    assert "indexed records only" in out["reply"]


@pytest.mark.parametrize("export", ["excel", "google_sheet"])
def test_real_search_snapshot_and_export_obey_window(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    export: str,
) -> None:
    """Real SQL excludes old rows before counting, snapshotting, and both XLSX paths."""
    path = _db(tmp_path)
    scope = recent_award_scope(BRETT_ASK, None, today=TODAY)
    assert scope is not None
    monkeypatch.setattr(
        google_sheets, "create_sheet", lambda *_args: ("unconfigured", "Unavailable")
    )
    text, artifact = search_leads(
        **scope.constrain(
            {
                "state": "CA",
                "org_type": "school",
                "export": export,
                "result_scope": "all",
                "db_path": path,
                "requester_slack": "U1",
                "workspace": "T1",
                "channel": "C1",
                "thread_ts": "1.0",
            }
        )
    )
    assert artifact is not None
    try:
        assert artifact.path.name == "grant_search_CA.xlsx"
        workbook = load_workbook(artifact.path, read_only=True)
        rows = list(workbook.active.values)
        workbook.close()
        assert len(rows) == 2
        assert rows[1][rows[0].index("entity_name")] == "Modesto City Schools"
        conn = db.connect(path)
        try:
            snapshot = conn.execute("SELECT * FROM search_requests").fetchone()
            assert snapshot is not None
            filters = json.loads(snapshot["filters_json"])
            assert filters["date_from"] == "2026-03-08"
            assert filters["date_to"] == "2026-09-08"
        finally:
            conn.close()
        assert "1" in text
    finally:
        artifact.cleanup()


@pytest.mark.parametrize("state", ["../../secret", "PA/ME", "", "California"])
def test_export_filename_does_not_accept_paths(state: str) -> None:
    """Only two-letter state labels may enter the filename helper."""
    assert search_export_filename(state) == "grant_search.xlsx"


def test_real_sql_boundary_future_unknown_and_unverified(tmp_path: Path) -> None:
    """Only evidenced dates on or inside both boundaries may reach recent results."""
    path = tmp_path / "bounds.db"
    conn = db.connect(path)
    try:
        for key, stamp in [
            ("before", "2026-03-07"),
            ("start", "2026-03-08"),
            ("end", "2026-09-08"),
            ("future", "2026-09-09"),
            ("unknown", "2026-04-01"),
            ("unverified", "2026-04-02"),
        ]:
            _insert(
                conn,
                "usaspending:16.071",
                key,
                f"{key} School",
                "PA",
                "SVPP",
                1000.0,
                stamp,
                "2028-01-01",
                LeadGrade.GOLD,
            )
        conn.execute(
            "UPDATE funding_events SET occurred_on=NULL WHERE occurred_on='2026-04-01'"
        )
        conn.execute(
            "UPDATE funding_events SET verification_status='needs-testing' WHERE occurred_on='2026-04-02'"
        )
        conn.commit()
    finally:
        conn.close()
    scope = recent_award_scope(BRETT_ASK, None, today=TODAY)
    assert scope is not None
    text, artifact = search_leads(
        **scope.constrain({"state": "PA", "org_type": "school", "db_path": path})
    )
    assert artifact is None
    assert "start School" in text and "end School" in text
    for excluded in ("before", "future", "unknown", "unverified"):
        assert f"{excluded} School" not in text


def test_with_contacts_and_exhaustion_keep_scope(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Paid contact followups cannot widen selection; the exhaustion answer stays dated."""
    executions: list[dict[str, object]] = []

    def capture(
        _name: str, args: dict[str, object], *_args: object, **_kwargs: object
    ) -> tuple[str, None]:
        """Fake contact work and capture actual filters without provider calls."""
        executions.append(args)
        return "No grants matched.", None

    monkeypatch.setattr(search_recency, "utc_today", lambda: TODAY)
    monkeypatch.setattr(conversation, "MAX_TOOL_TURNS", 1)
    monkeypatch.setattr(
        conversation,
        "Anthropic",
        _scripted_client(
            [
                [
                    {
                        "state": "PA",
                        "with_contacts": True,
                        "limit": 5,
                    }
                ]
            ]
        ),
    )
    monkeypatch.setattr(tools, "run_tool", capture)
    out = conversation.respond(
        "find contacts for the top 5", None, ["rep: " + BRETT_ASK]
    )
    assert executions[0]["date_from"] == "2026-03-08"
    assert "2026-03-08 through 2026-09-08" in out["reply"]


def test_conversation_dispatch_to_real_sql_snapshot_and_workbook(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """An omitted date in the model call becomes persisted filters and real XLSX rows."""
    path = _db(tmp_path)
    monkeypatch.setattr(db, "DEFAULT_DB_PATH", path)
    monkeypatch.setattr(search_recency, "utc_today", lambda: TODAY)
    monkeypatch.setattr(
        conversation,
        "Anthropic",
        _scripted_client(
            [
                [
                    {
                        "state": "CA",
                        "org_type": "school",
                        "export": "excel",
                        "result_scope": "all",
                    }
                ]
            ]
        ),
    )
    out = conversation.respond(
        "all California schools awarded grants recently in Excel",
        None,
        requester_slack="U1",
        workspace="T1",
        channel="C1",
        thread_ts="1.0",
    )
    assert len(out["files"]) == 1
    artifact = out["files"][0]
    try:
        workbook = load_workbook(artifact.path, read_only=True)
        rows = list(workbook.active.values)
        workbook.close()
        assert len(rows) == 2
        assert rows[1][rows[0].index("entity_name")] == "Modesto City Schools"
        conn = db.connect(path)
        try:
            snapshot = conn.execute("SELECT * FROM search_requests").fetchone()
            assert snapshot is not None
            assert json.loads(snapshot["filters_json"])["date_from"] == "2026-03-08"
            assert len(json.loads(snapshot["result_lead_ids_json"])) == 1
        finally:
            conn.close()
    finally:
        artifact.cleanup()
